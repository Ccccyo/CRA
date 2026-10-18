"""Mock CRA Excel（演示用，结构从简）：AI 写黄色格子 → 助手内查看/编辑 → 模拟 Run Macro 回填结果。

同一份文件来回走：Excel 宏从 SAP 拉完数据 → 上传到助手 → 助手在同一份文件的黄色格子里填 AI 结果 → 下载回 Excel。

工作表：
- Financial Statements：三年开关（C2:E2）+ 本年（AI 抽取，待确认）+ 往年两列（宏从 SAP 导出）
- Executive Summary：BP、客户、申请额度；定性评分、额度方法、意见；Run Macro 的结果回填在这里
- Header params：客户地址、分析师填写项、SAP 应收的付款记录、导出信息（键 / 值）
- TCURR：汇率（宏从 SAP 表 TCURR 导出）
- AI Import：AI 填写后才有，每个数字的出处、模型自评、校验提示（审计用；VBA Import AI 读这张表）
保留与真实文件同名的命名单元格（BP、ENDDATE、PayHabits、Score…），VBA 宏按这些名字读值。
样式只用一种淡黄色标「待人工确认」，有校验提示的格子用红字 + 批注。
"""
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from cra.casefile import Case
from cra.config import ROOT
from cra.draft.narrative import QUALITATIVE_OPTIONS
from cra.extract.schema import FIELD_LABELS, KEY_FIELDS, ExtractedValue
from cra.validate import Flag
from cra.validate.financial_checks import Values

INPUT = PatternFill("solid", fgColor="FFF9DB")  # 唯一的底色：待人工确认的输入格
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
RED = Font(color="C00000")
WRAP = Alignment(wrap_text=True, vertical="top")
UNDERLINE = Border(bottom=Side(style="thin", color="808080"))

YEAR_OPTIONS = ["Yes - Bring in Prior Year Financials", "No - Enter Manually", "No Action Needed"]
CL_METHODS = ["A. % Tangible Net Worth", "C. % Net Worth"]
YEAR_COLS = ["C", "D", "E"]  # Current Year / Prior Year / Prior Year - 1

FS_SHEET, ES_SHEET, AI_SHEET = "Financial Statements", "Executive Summary", "AI Import"
HEADER_SHEET, TCURR_SHEET = "Header params", "TCURR"
TEMPLATE_VERSION = "CRA-2026.1"
# Header params 的键 → 说明
HEADER_KEYS = {
    "TemplateVersion": "模板版本", "ExportedAt": "宏导出时间", "ExportedBy": "导出人",
    "CustomerName": "客户名称（SAP 主数据）", "Country": "国家代码", "State": "州 / 省", "City": "城市",
    "Business": "业务说明（分析师填写）", "CompanyTypeHint": "公司类型线索（分析师填写）",
    "CustomerDistinction": "战略重要性（商务部门）", "SovereignHint": "国家风险线索（可空，助手会查 OECD）",
    "PayHabitsFromAR": "付款习惯档位（SAP 应收）", "PayHabitsDetail": "付款记录说明（SAP 应收）",
}
FIRST_FIELD_ROW = 10
FIELD_ROW = {code: FIRST_FIELD_ROW + i for i, code in enumerate(KEY_FIELDS)}
BALANCE_ROW = FIRST_FIELD_ROW + len(KEY_FIELDS) + 1

# Executive Summary 版面
QUAL_ROW = {"pay_habits": 8, "company_type": 9, "sovereign_rating": 10, "customer_distinction": 11}
QUAL_LABEL = {"pay_habits": "Twelve Months Pay Habits", "company_type": "Company Type",
              "sovereign_rating": "Sovereign Rating", "customer_distinction": "Customer Distinction"}
QUAL_NAME = {"pay_habits": "PayHabits", "company_type": "CompanyType", "sovereign_rating": "SovereignRating",
             "customer_distinction": "CustomerDistinction"}
RATING_ROW = 12
DUNS_ROW = 13
RESULT_ROWS = {"financial": 15, "qualitative": 16, "overall": 17, "risk_class": 18, "method": 19,
               "suggested_cl": 20, "approval": 21}
DESC_ROW, RISK_ROW = 24, 27


def mock_path(case_id: Optional[str] = None) -> Path:
    """Mock Excel 放在项目根目录（D&B 文件夹）。"""
    return ROOT / (f"CRA_Mock_{case_id}.xlsx" if case_id else "CRA_Mock_Template.xlsx")


def _name(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    col, row = cell.rstrip("0123456789"), cell[len(cell.rstrip("0123456789")):]
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!${col}${row}")


def _dropdown(ws, cell: str, ref: str) -> None:
    dv = DataValidation(type="list", formula1=f"={ref}")
    ws.add_data_validation(dv)
    dv.add(ws[cell])


def _lists(wb: Workbook) -> Dict[str, str]:
    """隐藏的 Lists 表放下拉选项，返回 {名称: 引用}。"""
    ws = wb.create_sheet("Lists")
    ws.sheet_state = "hidden"
    groups = {"year": YEAR_OPTIONS, "method": CL_METHODS, **QUALITATIVE_OPTIONS}
    refs = {}
    for c, (key, opts) in enumerate(groups.items(), 1):
        letter = chr(ord("A") + c - 1)
        ws.cell(row=1, column=c, value=key)
        for r, v in enumerate(opts, 2):
            ws.cell(row=r, column=c, value=v)
        refs[key] = f"Lists!${letter}$2:${letter}${len(opts) + 1}"
    return refs


def _financial_statements(wb: Workbook, refs: Dict[str, str], case: Optional[Case], periods: List[str]) -> None:
    ws = wb.create_sheet(FS_SHEET)
    ws["A1"], ws["A1"].font = "Financial Statements（MOCK）", TITLE
    ws["A2"] = "Export Existing Financials from SAP?"
    defaults = ["No - Enter Manually"] + ["Yes - Bring in Prior Year Financials"] * 2
    for col, default in zip(YEAR_COLS, defaults):
        ws[f"{col}2"] = default
        _dropdown(ws, f"{col}2", refs["year"])
    labels = {3: "Bus. Partner", 4: "Crd. Segment", 5: "St. End Date", 6: "Stmt Currency", 7: "来源"}
    for row, label in labels.items():
        ws[f"A{row}"] = label
    if case:
        ws["C3"], ws["C4"] = case.bp, case.credit_segment
    _name(wb, "BP", FS_SHEET, "C3")
    _name(wb, "CreditSegment", FS_SHEET, "C4")
    for i, col in enumerate(YEAR_COLS):
        suffix = "" if i == 0 else str(i + 1)
        _name(wb, f"ENDDATE{suffix}", FS_SHEET, f"{col}5")
        _name(wb, f"StmtCurrency{suffix}", FS_SHEET, f"{col}6")
        if i < len(periods):
            ws[f"{col}5"] = periods[i]
            ws[f"{col}6"] = case.statement_currency if case else ""
            ws[f"{col}7"] = "AI 抽取（待确认）" if i == 0 else "SAP 导出（宏）"

    for c, h in enumerate(["科目", "field_code", "Current Year", "Prior Year", "Prior Year - 1"], 1):
        cell = ws.cell(row=FIRST_FIELD_ROW - 1, column=c, value=h)
        cell.font, cell.border = BOLD, UNDERLINE
    for code, row in FIELD_ROW.items():
        ws[f"A{row}"] = FIELD_LABELS[code]
        ws[f"B{row}"] = code
        ws[f"B{row}"].font = Font(color="808080", size=9)
        for col in YEAR_COLS:
            ws[f"{col}{row}"].number_format = "#,##0"
    ta, tl, te = FIELD_ROW["total_assets"], FIELD_ROW["total_liabilities"], FIELD_ROW["total_equity"]
    ws[f"A{BALANCE_ROW}"], ws[f"A{BALANCE_ROW}"].font = "Balance Sheet Check", BOLD
    for col in YEAR_COLS:
        ws[f"{col}{BALANCE_ROW}"] = (f'=IF({col}{ta}="","",IF(ABS({col}{ta}-({col}{tl}+{col}{te}))'
                                     f'<=0.005*ABS({col}{ta}),"OK","CHECK"))')
    for col, w in zip("ABCDE", (22, 26, 18, 18, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"C{FIRST_FIELD_ROW}"


def _executive_summary(wb: Workbook, refs: Dict[str, str], case: Optional[Case]) -> None:
    ws = wb.create_sheet(ES_SHEET)
    ws["A1"], ws["A1"].font = "Executive Summary（MOCK）", TITLE
    ws["A3"], ws["B3"] = "Bus. Partner", case.bp if case else None
    ws["A4"], ws["B4"] = "客户", case.customer["name"] if case else None
    ws["A5"], ws["B5"] = "Requested CL (USD)", case.requested_cl_usd if case else None
    ws["B5"].number_format = "#,##0"

    ws["A7"], ws["A7"].font = "定性评分（分析师选择）", BOLD
    ws["C7"], ws["C7"].font = "AI 建议 / 理由", BOLD
    for key, row in QUAL_ROW.items():
        ws[f"A{row}"] = QUAL_LABEL[key]
        _dropdown(ws, f"B{row}", refs[key])
        ws[f"C{row}"].alignment = WRAP
        _name(wb, QUAL_NAME[key], ES_SHEET, f"B{row}")
    ws[f"A{RATING_ROW}"] = "Corporate Credit Risk Rating"
    ws[f"C{RATING_ROW}"] = "按 D&B 评级查表，不由 AI 判断"
    _name(wb, "CorporateCreditRiskRating", ES_SHEET, f"B{RATING_ROW}")
    ws[f"A{DUNS_ROW}"] = "D&B DUNS（助手里已 Merge）"
    ws[f"C{DUNS_ROW}"] = "Import AI 时一并写入，供 VBA 核对 eRAM 绑定"
    _name(wb, "DUNS", ES_SHEET, f"B{DUNS_ROW}")

    ws["A14"], ws["A14"].font = "Run Macro 结果（演示中为模拟，真实流程由 SAP 回填）", BOLD
    titles = {"financial": "Financial Score (70%)", "qualitative": "Qualitative Score (30%)",
              "overall": "Overall Risk Rating Score", "risk_class": "Risk Category",
              "method": "Credit Limit Calculation Method", "suggested_cl": "Suggested Credit Limit (USD)",
              "approval": "审批层级"}
    for key, row in RESULT_ROWS.items():
        ws[f"A{row}"] = titles[key]
    ws[f"B{RESULT_ROWS['method']}"] = CL_METHODS[0]
    _dropdown(ws, f"B{RESULT_ROWS['method']}", refs["method"])
    ws[f"B{RESULT_ROWS['suggested_cl']}"].number_format = "#,##0"
    _name(wb, "Score", ES_SHEET, f"B{RESULT_ROWS['overall']}")

    for row, title in ((DESC_ROW - 1, "Company Description"), (RISK_ROW - 1, "Risk Comments - Used for SCASE")):
        ws[f"A{row}"], ws[f"A{row}"].font = title, BOLD
        ws.merge_cells(f"A{row + 1}:C{row + 1}")
        ws[f"A{row + 1}"].alignment = WRAP
        ws.row_dimensions[row + 1].height = 110 if row + 1 == RISK_ROW else 50
    _name(wb, "CompanyDescription", ES_SHEET, f"A{DESC_ROW}")
    _name(wb, "RiskComments", ES_SHEET, f"A{RISK_ROW}")
    for col, w in zip("ABC", (34, 34, 70)):
        ws.column_dimensions[col].width = w


def _header_params(wb: Workbook, case: Optional[Case], exported_by: str) -> None:
    ws = wb.create_sheet(HEADER_SHEET)
    ws.append(["Key", "Value", "说明"])
    for c in ws[1]:
        c.font, c.border = BOLD, UNDERLINE
    i = case.inputs if case else {}
    values = {"TemplateVersion": TEMPLATE_VERSION}
    if case:
        values.update({
            "ExportedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "ExportedBy": exported_by,
            "CustomerName": case.customer["name"], "Country": case.customer["country"],
            "State": case.customer["state"], "City": case.customer["city"], "Business": case.business,
            "CompanyTypeHint": i.get("company_type_hint", ""), "CustomerDistinction": i.get("customer_distinction_from_sales", ""),
            "SovereignHint": i.get("sovereign_hint", ""), "PayHabitsFromAR": i.get("pay_habits_from_ar", ""),
            "PayHabitsDetail": i.get("pay_habits_detail", ""),
        })
    for k, note in HEADER_KEYS.items():
        ws.append([k, values.get(k, ""), note])
    for col, w in zip("ABC", (22, 44, 36)):
        ws.column_dimensions[col].width = w


def _tcurr(wb: Workbook, fx: Optional[Dict]) -> None:
    ws = wb.create_sheet(TCURR_SHEET)
    ws.append(["currency", "valid_date", "units_per_usd"])
    for c in ws[1]:
        c.font, c.border = BOLD, UNDERLINE
    for (ccy, d), rate in sorted((fx or {}).items()):
        ws.append([ccy, d, rate])


def build_mock_workbook(case: Optional[Case] = None, sap_values: Optional[Values] = None,
                        periods: Optional[List[str]] = None, switches: Optional[List[str]] = None,
                        fx: Optional[Dict] = None, exported_by: str = "analyst@example.com") -> Workbook:
    """空白模板，或「宏已从 SAP 拉完数据」的工作簿：抬头、三年开关、往年财报、客户信息、汇率都填好，本年度留空等 AI 填。"""
    wb = Workbook()
    wb.remove(wb.active)
    refs = _lists(wb)
    periods = periods or ([case.new_period] + sorted(sap_values or {}, reverse=True)[:2] if case else [])
    _financial_statements(wb, refs, case, periods)
    _executive_summary(wb, refs, case)
    _header_params(wb, case, exported_by)
    _tcurr(wb, fx)
    if switches:
        for col, sw in zip(YEAR_COLS, switches):
            wb[FS_SHEET][f"{col}2"] = sw
    wb.move_sheet(wb["Lists"], offset=len(wb.sheetnames) - 1)  # 隐藏的 Lists 放最后
    if sap_values:
        ws = wb[FS_SHEET]
        for col, period in zip(YEAR_COLS[1:], periods[1:]):
            for code, row in FIELD_ROW.items():
                ws[f"{col}{row}"] = sap_values.get(period, {}).get(code)
    return wb


def fill_ai_results(wb: Workbook, extracted: List[ExtractedValue], flags: List[Flag], new_period: str,
                    qualitative: Dict[str, Dict[str, str]], company_description: str, risk_comments: str,
                    dnb_rating: str, model_name: str, dnb_duns: str = "") -> None:
    """AI 结果写进待确认格子（淡黄）；有校验提示的红字 + 批注；并生成 AI Import 审计表。"""
    ws = wb[FS_SHEET]
    flagged = {f.field: f for f in flags}
    by_field = {v.field: v for v in extracted if v.period_end == new_period}
    for code, row in FIELD_ROW.items():
        cell = ws[f"C{row}"]
        v = by_field.get(code)
        cell.value, cell.fill = (v.normalized if v else None), INPUT
        notes = []
        if v is not None:
            notes.append(f"第 {v.page} 页：{v.source_text or '（无原文）'}（模型自评 {v.confidence}）")
            if v.note:
                notes.append(v.note)
        if code in flagged:
            notes.append("⚠ " + flagged[code].message)
            cell.font = RED
        if notes:
            cell.comment = Comment("\n".join(notes), "AI 助手", width=300, height=100)

    es = wb[ES_SHEET]
    for key, row in QUAL_ROW.items():
        q = qualitative.get(key, {})
        es[f"B{row}"], es[f"B{row}"].fill = q.get("final"), INPUT
        es[f"C{row}"] = f"{q['ai']}：{q.get('reason', '')}" if q.get("ai") else None
    es[f"B{RATING_ROW}"] = dnb_rating
    es[f"B{DUNS_ROW}"] = dnb_duns or None
    es[f"A{DESC_ROW}"], es[f"A{RISK_ROW}"] = company_description, risk_comments
    es[f"A{DESC_ROW}"].fill = es[f"A{RISK_ROW}"].fill = INPUT

    if AI_SHEET in wb.sheetnames:  # 同一份文件再填一次：换掉上一版的审计表
        del wb[AI_SHEET]
    ai = wb.create_sheet(AI_SHEET, index=len(wb.sheetnames) - 1)
    headers = ["field_code", "科目", "期间", "数值", "页码", "原文", "模型自评", "校验提示", "写入位置"]
    for c, h in enumerate(headers, 1):
        cell = ai.cell(row=1, column=c, value=h)
        cell.font, cell.border = BOLD, UNDERLINE
    for v in extracted:
        f = flagged.get(v.field)
        target = f"'{FS_SHEET}'!C{FIELD_ROW[v.field]}" if v.field in FIELD_ROW and v.period_end == new_period else ""
        ai.append([v.field, FIELD_LABELS.get(v.field, v.field), v.period_end, v.normalized, v.page,
                   v.source_text, v.confidence, f.message if f else "", target])
    ai.append([])
    ai.append([f"模型：{model_name} · 生成时间：{datetime.now():%Y-%m-%d %H:%M}"])
    for col, w in zip("ABCDEFGHI", (24, 14, 12, 16, 6, 48, 8, 40, 30)):
        ai.column_dimensions[col].width = w
    ai.freeze_panes = "A2"


# ---------- 在助手里读写这份 Excel ----------

def read_workbook(path: Path) -> Dict:
    """读出页面需要的内容：三年数字、定性评分、额度方法、意见、Run Macro 结果。"""
    wb = load_workbook(path)
    fs, es = wb[FS_SHEET], wb[ES_SHEET]
    periods = [fs[f"{c}5"].value for c in YEAR_COLS]
    values: Values = {}
    for col, period in zip(YEAR_COLS, periods):
        if period:
            values[str(period)] = {code: fs[f"{col}{row}"].value for code, row in FIELD_ROW.items()
                                   if fs[f"{col}{row}"].value is not None}
    comments = {code: fs[f"C{row}"].comment.text for code, row in FIELD_ROW.items() if fs[f"C{row}"].comment}
    return {
        "periods": [str(p) for p in periods if p], "values": values, "comments": comments,
        "bp": fs["C3"].value, "currency": fs["C6"].value,
        "qual": {k: es[f"B{r}"].value for k, r in QUAL_ROW.items()},
        "qual_ai": {k: es[f"C{r}"].value for k, r in QUAL_ROW.items()},
        "rating": es[f"B{RATING_ROW}"].value, "duns": es[f"B{DUNS_ROW}"].value, "requested": es["B5"].value,
        "method": es[f"B{RESULT_ROWS['method']}"].value,
        "results": {k: es[f"B{r}"].value for k, r in RESULT_ROWS.items() if k != "method"},
        "desc": es[f"A{DESC_ROW}"].value or "", "risk": es[f"A{RISK_ROW}"].value or "",
        "ai_rows": [[c.value for c in row] for row in wb[AI_SHEET].iter_rows(min_row=1)] if AI_SHEET in wb.sheetnames else [],
    }


def save_edits(path: Path, current: Dict[str, Optional[float]], qual: Dict[str, Optional[str]], method: str,
               desc: str, risk: str) -> None:
    """把助手里的修改写回 Excel（只改输入格子，不碰往年 SAP 列和公式）。"""
    wb = load_workbook(path)
    fs, es = wb[FS_SHEET], wb[ES_SHEET]
    for code, val in current.items():
        fs[f"C{FIELD_ROW[code]}"] = val
    for key, val in qual.items():
        es[f"B{QUAL_ROW[key]}"] = val
    es[f"B{RESULT_ROWS['method']}"] = method
    es[f"A{DESC_ROW}"], es[f"A{RISK_ROW}"] = desc, risk
    wb.save(path)


def write_risk_opinion(
    path: Path,
    company_description: str,
    risk_comments: str,
    qualitative: Optional[Dict[str, Dict[str, str]]] = None,
    dnb_rating: str = "",
    dnb_duns: str = "",
) -> None:
    """Write the approved risk review fields, preserving the supplied CRA template."""
    wb = load_workbook(path)
    if ES_SHEET not in wb.sheetnames:
        raise ValueError(f"工作簿缺少 {ES_SHEET} 工作表")
    es = wb[ES_SHEET]
    es[f"A{DESC_ROW}"] = company_description
    es[f"A{RISK_ROW}"] = risk_comments
    es[f"A{DESC_ROW}"].fill = INPUT
    es[f"A{RISK_ROW}"].fill = INPUT
    es[f"A{DESC_ROW}"].alignment = WRAP
    es[f"A{RISK_ROW}"].alignment = WRAP
    for key, row in QUAL_ROW.items():
        item = (qualitative or {}).get(key, {})
        choice = item.get("choice") or item.get("final")
        if choice:
            es[f"B{row}"] = choice
            es[f"B{row}"].fill = INPUT
        reason = item.get("reason")
        if reason:
            es[f"C{row}"] = reason
            es[f"C{row}"].alignment = WRAP
    if dnb_rating:
        es[f"B{RATING_ROW}"] = dnb_rating
    if dnb_duns:
        es[f"B{DUNS_ROW}"] = dnb_duns
    wb.save(path)


def save_macro_results(path: Path, results: Dict[str, object]) -> None:
    wb = load_workbook(path)
    es = wb[ES_SHEET]
    for key, val in results.items():
        if key in RESULT_ROWS and key != "method":
            es[f"B{RESULT_ROWS[key]}"] = val
    wb.save(path)
