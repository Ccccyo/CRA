"""AI 交接包 = CRA Excel 文件本身：导入和导出是同一个格式，只是多填了内容。

动线：Excel 填 BP → VBA 从 SAP 拉往年财报、汇率、付款记录（填进模板）→ 把这份 CRA 文件上传到助手
     → 助手里认公司（Lookup + Merge）、读新财报、写意见 → 在**同一份文件**的黄色格子里填好 → 下载
     → Excel 里 Import AI（或直接用这份文件）→ Run Macro → SAP

助手从这几处读（格式见 excel_mock.py）：
- Financial Statements：BP（C3）、信用段（C4）、三年开关（C2:E2）、期末（C5:E5）、币种（C6）、往年财报（D、E 列）
- Executive Summary：申请额度（B5）
- Header params：客户地址、分析师填写项、SAP 应收的付款记录、导出信息
- TCURR：汇率
"""
import json
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from cra.excel_mock import (ES_SHEET, FIELD_ROW, FS_SHEET, HEADER_SHEET, TCURR_SHEET, YEAR_COLS, YEAR_OPTIONS,
                            build_mock_workbook)

SWITCH_YES, SWITCH_MANUAL, SWITCH_NONE = YEAR_OPTIONS
SWITCHES = YEAR_OPTIONS


@dataclass
class Package:
    meta: Dict[str, str]
    case: Dict[str, str]
    years: List[Tuple[str, str]]  # (期末日期, 开关)
    sap: Dict[str, Dict[str, float]]  # {期间: {科目: 数值}}
    fx: Dict[Tuple[str, str], float]  # {(币种, 日期): 1 美元兑多少本币}
    ar: Dict[str, str]
    errors: List[str] = field(default_factory=list)

    @property
    def manual_years(self) -> List[str]:
        """需要手工补（由 AI 从客户财报抽取）的年份。"""
        return [y for y, sw in self.years if sw == SWITCH_MANUAL]

    @property
    def sap_years(self) -> List[str]:
        return [y for y, sw in self.years if sw == SWITCH_YES]


# ---------- 写（演示用：生成「宏已从 SAP 拉完数据」的 CRA 文件；真实场景由 VBA 填） ----------

def write_package(path: Path, case: Dict[str, object], years: List[Tuple[str, str]],
                  sap: Dict[str, Dict[str, float]], fx: Dict[Tuple[str, str], float], ar: Dict[str, str],
                  exported_by: str = "analyst@example.com") -> Path:
    """case 用交接包的键（BP、CustomerName、Country…）；years 按 C、D、E 列顺序（本年在前）。"""
    from cra.casefile import Case

    c = Case(case_id="", title="", bp=str(case["BP"]), credit_segment=str(case.get("CreditSegment") or "0001"),
             requested_cl_usd=float(case.get("RequestedCL_USD") or 0),
             statement_currency=str(case["StatementCurrency"]).upper(),
             customer={"name": case.get("CustomerName", ""), "country": case.get("Country", ""),
                       "state": case.get("State", ""), "city": case.get("City", "")},
             expected_entity_id="", business=case.get("Business", ""), new_year={"period_end": years[0][0]},
             inputs={"company_type_hint": case.get("CompanyTypeHint", ""),
                     "customer_distinction_from_sales": case.get("CustomerDistinction", ""),
                     "sovereign_hint": case.get("SovereignHint", ""), "pay_habits_from_ar": ar.get("PayHabits", ""),
                     "pay_habits_detail": ar.get("PayHabitsDetail", "")})
    wb = build_mock_workbook(c, sap, periods=[y for y, _ in years], switches=[sw for _, sw in years], fx=fx,
                             exported_by=exported_by)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------- 读 ----------

def _kv(ws) -> Dict[str, str]:
    return {str(r[0]).strip(): ("" if r[1] is None else str(r[1]).strip())
            for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}


def _text(v) -> str:
    return "" if v is None else str(v).strip()


@lru_cache(maxsize=32)
def _read_cached(path: str, mtime: float) -> Package:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    missing = [s for s in (FS_SHEET, ES_SHEET, HEADER_SHEET) if s not in wb.sheetnames]
    if missing:
        return Package({}, {}, [], {}, {}, {}, [f"缺少工作表：{'、'.join(missing)}（请上传 CRA 模板格式的文件）"])
    fs, es, hp = wb[FS_SHEET], wb[ES_SHEET], _kv(wb[HEADER_SHEET])
    meta = {k: hp.get(k, "") for k in ("TemplateVersion", "ExportedAt", "ExportedBy")}
    case = {"BP": _text(fs["C3"].value), "CreditSegment": _text(fs["C4"].value),
            "StatementCurrency": _text(fs["C6"].value) or _text(fs["D6"].value),
            "RequestedCL_USD": _text(es["B5"].value),
            "CustomerName": hp.get("CustomerName") or _text(es["B4"].value),
            **{k: hp.get(k, "") for k in ("Country", "State", "City", "Business", "CompanyTypeHint",
                                         "CustomerDistinction", "SovereignHint")}}
    years, sap = [], {}
    for col in YEAR_COLS:
        end, sw = _text(fs[f"{col}5"].value)[:10], _text(fs[f"{col}2"].value)
        if not end:
            continue
        years.append((end, sw))
        if sw == SWITCH_YES:
            vals = {code: float(fs[f"{col}{row}"].value) for code, row in FIELD_ROW.items()
                    if fs[f"{col}{row}"].value not in (None, "")}
            if vals:
                sap[end] = vals
    fx = {}
    if TCURR_SHEET in wb.sheetnames:
        fx = {(str(r[0]), str(r[1])[:10]): float(r[2])
              for r in wb[TCURR_SHEET].iter_rows(min_row=2, values_only=True) if r[0] and r[2] is not None}
    ar = {"PayHabits": hp.get("PayHabitsFromAR", ""), "PayHabitsDetail": hp.get("PayHabitsDetail", "")}

    errors = []
    if not years:
        return Package(meta, case, years, sap, fx, ar, ["这是空白模板：期末日期（C5:E5）、BP 都还没填。"
                                                       "请先在 Excel 里填好，并运行宏把 SAP 数据拉进来"])
    if not case["BP"]:
        errors.append("Financial Statements 的 BP（C3）是空的")
    if not case["StatementCurrency"]:
        errors.append("Financial Statements 的报表币种（C6）是空的")
    bad = [sw for _, sw in years if sw not in SWITCHES]
    if bad:
        errors.append(f"三年开关（C2:E2）有无效值：{bad}")
    manual = [y for y, sw in years if sw == SWITCH_MANUAL]
    if not manual:
        errors.append("没有标为「No - Enter Manually」的年份：没有需要 AI 补的财报")
    elif _text(fs["C2"].value) != SWITCH_MANUAL:
        errors.append("要 AI 补的年份必须在本年度列（C 列）：请把 C2 设为「No - Enter Manually」")
    for y, sw in years:
        if sw == SWITCH_YES and y not in sap:
            errors.append(f"{y} 标为从 SAP 导出，但这一列没有数字：先在 Excel 里运行宏把 SAP 数据拉进来")
    return Package(meta, case, years, sap, fx, ar, errors)


def read_package(path) -> Package:
    p = Path(path)
    return _read_cached(str(p), p.stat().st_mtime)


# ---------- 导入：建案或更新案件 ----------

def _data_dir() -> Path:
    from cra.casefile import _data_dir as d
    return d()


def import_package(src: Path, force_new: bool = False):
    """读交接包 → 保存一份副本 → 按 BP 找到已有案件就更新，没有就新建。返回 (案件, 是否新建, 错误列表)。"""
    from cra.casefile import Case, load_cases, next_case_id
    from cra.history import load_history

    pkg = read_package(src)
    if pkg.errors:
        return None, False, pkg.errors
    c = pkg.case
    bp = c["BP"]
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = _data_dir() / "handoff" / "imported" / f"{bp}_{stamp}.xlsx"
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(src, dest)
    rel = str(dest.relative_to(_data_dir().parent))

    updates = dict(
        bp=bp, credit_segment=c.get("CreditSegment") or "0001",
        requested_cl_usd=float(c.get("RequestedCL_USD") or 0), statement_currency=c["StatementCurrency"].upper(),
        customer={"name": c.get("CustomerName", ""), "country": c.get("Country", "").upper(),
                  "state": c.get("State", ""), "city": c.get("City", "")},
        inputs={"pay_habits_from_ar": pkg.ar.get("PayHabits") or "No history",
                "pay_habits_detail": pkg.ar.get("PayHabitsDetail") or "（交接包未提供）",
                "customer_distinction_from_sales": c.get("CustomerDistinction") or "（待补充）",
                "company_type_hint": c.get("CompanyTypeHint") or "（待补充）",
                "sovereign_hint": c.get("SovereignHint") or "（助手查 OECD 国家风险）"},
        package=rel,
    )
    existing = None if force_new else next((x for x in load_cases() if x.bp == bp), None)
    if existing:
        data = {**existing.__dict__, **updates}
        data["new_year"] = {**existing.new_year, "period_end": pkg.manual_years[0]}
        if c.get("Business"):
            data["business"] = c["Business"]
        case, created = Case(**data), False
    else:
        past = [r for r in load_history() if r["bp"] == bp]
        title = f"{c.get('CustomerName')} {'年度复审' if past else '新客户授信'}"
        case = Case(case_id=next_case_id(), title=title, expected_entity_id="", business=c.get("Business") or "待补充",
                    new_year={"period_end": pkg.manual_years[0], "pdf": "（待上传）", "scan": False}, **updates)
        created = True
    path = _data_dir() / "cases" / f"{case.case_id}.json"
    path.write_text(json.dumps(case.__dict__, ensure_ascii=False, indent=2), encoding="utf-8")
    return case, created, []


def case_package(case) -> Optional[Package]:
    if not getattr(case, "package", None):
        return None
    p = _data_dir().parent / case.package
    return read_package(p) if p.exists() else None


def case_sap(case) -> Dict[str, Dict[str, float]]:
    """案件的往年财报：只来自导入的交接包（没导入就是空的，AI 不会去后台找 SAP 数据）。"""
    pkg = case_package(case)
    return {y: pkg.sap[y] for y in pkg.sap_years if y in pkg.sap} if pkg else {}


def case_fx(case) -> Dict[Tuple[str, str], float]:
    pkg = case_package(case)
    return pkg.fx if pkg else {}


def sample_path(case_id: str) -> Path:
    return _data_dir() / "handoff" / "samples" / f"{case_id}.xlsx"
