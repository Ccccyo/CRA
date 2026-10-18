"""Mock CRA Excel：生成 → 助手内读取 → 修改保存 → 模拟 Run Macro 回填，全程往返。"""
import json

import pytest
from openpyxl import load_workbook

from cra.casefile import load_cases, load_sample, load_sap_export
from cra.excel_mock import (FIELD_ROW, FS_SHEET, RED, build_mock_workbook, fill_ai_results, read_workbook,
                            save_edits, save_macro_results)
from cra.extract.extractor import extract_from_pages
from cra.ingest.pdf_reader import Page
from cra.llm import FakeLLM
from cra.mock_sap import run_macro

QUAL = {"pay_habits": {"final": "Prompt"}, "company_type": {"final": "Private - audited"},
        "sovereign_rating": {"final": "grade A or above"}, "customer_distinction": {"final": "Strategic"}}


def _write(tmp_path, case_id):
    case = next(c for c in load_cases() if c.case_id == case_id)
    s = load_sample("extraction", case_id)
    rep = extract_from_pages([Page(p["number"], p["text"]) for p in s["pages"]], FakeLLM(json.dumps(s["llm_output"])))
    wb = build_mock_workbook(case, load_sap_export(case.bp))
    fill_ai_results(wb, rep.result.values, rep.flags, case.new_period, QUAL, "desc", "risk", "3A2", "fake")
    path = tmp_path / f"{case_id}.xlsx"
    wb.save(path)
    return path


def test_layout_is_simple_and_keeps_vba_names():
    wb = build_mock_workbook()
    assert [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"] == [FS_SHEET, "Executive Summary", "Header params", "TCURR"]
    for name in ["BP", "CreditSegment", "ENDDATE", "ENDDATE2", "StmtCurrency", "PayHabits", "CompanyType",
                 "SovereignRating", "CustomerDistinction", "CorporateCreditRiskRating", "Score"]:
        assert name in wb.defined_names, name


def test_ai_current_year_and_sap_prior_years(tmp_path):
    data = read_workbook(_write(tmp_path, "CASE-001"))
    assert data["periods"] == ["2025-12-31", "2024-12-31", "2023-12-31"]
    assert [data["values"][p]["total_assets"] for p in data["periods"]] == [42_000_000, 39_500_000, 37_000_000]
    assert "第 41 页" in data["comments"]["total_assets"]


def test_flagged_cells_are_red_with_comment(tmp_path):
    ws = load_workbook(_write(tmp_path, "CASE-003"))[FS_SHEET]
    cell = ws[f"C{FIELD_ROW['inventories']}"]
    assert cell.font.color.rgb == RED.color.rgb and "⚠" in cell.comment.text


def test_edit_then_macro_round_trip(tmp_path):
    path = _write(tmp_path, "CASE-003")
    data = read_workbook(path)
    current = dict(data["values"]["2025-12-31"], inventories=241_000_000.0)  # 人工按原文更正
    save_edits(path, current, {k: q["final"] for k, q in QUAL.items()}, "A. % Tangible Net Worth", "d", "r")
    data = read_workbook(path)
    assert data["values"]["2025-12-31"]["inventories"] == 241_000_000
    result = run_macro(data["values"]["2025-12-31"], data["qual"], data["rating"], data["method"], 800_000, 18.4)
    save_macro_results(path, result.as_cells())
    assert read_workbook(path)["results"]["risk_class"] == result.risk_class
    assert 1 <= result.overall <= 6 and result.suggested_cl >= 0


def test_macro_approval_escalates_when_request_exceeds_suggestion():
    v = {"operating_income": 5e6, "interest_expense": 6e5, "operating_cash_flow": 3.6e6, "total_liabilities": 2.46e7,
         "total_equity": 1.74e7, "net_sales": 6.8e7, "net_income": 3.3e6, "total_assets": 4.2e7, "intangible_assets": 1.9e6}
    small = run_macro(v, {k: q["final"] for k, q in QUAL.items()}, "3A2", "A. % Tangible Net Worth", 1_000)
    big = run_macro(v, {k: q["final"] for k, q in QUAL.items()}, "3A2", "A. % Tangible Net Worth", 50_000_000)
    assert "提高一级" not in small.approval and "提高一级" in big.approval
